# -*- coding: utf-8 -*-
"""Генерирует Word-документ: раздел 2 по образцу записки Раубы, данные BuildFlow."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


OUT = Path(r"C:\Users\Даниил\Downloads\Telegram Desktop\BuildFlow_section2_project_design.docx")


def set_doc_defaults(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def disable_mirror_margins(doc: Document):
    """Отключает зеркальные поля (часто используют для двусторонней вёрстки)."""
    settings_el = doc.settings.element
    for el in settings_el.findall(qn("w:mirrorMargins")):
        settings_el.remove(el)


def add_heading_center(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    if level <= 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_para(doc: Document, text: str, justify=True):
    p = doc.add_paragraph(text)
    if justify:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_table_caption(doc: Document, caption: str):
    p = doc.add_paragraph(caption)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.runs[0]
    r.bold = True


def fill_table(doc: Document, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        row_cells = tbl.rows[ri + 1].cells
        for ci, val in enumerate(row):
            row_cells[ci].text = val
    return tbl


def main():
    doc = Document()

    sect = doc.sections[0]
    sect.page_height = Mm(297)
    sect.page_width = Mm(210)
    sect.left_margin = Mm(30)
    sect.right_margin = Mm(15)
    sect.top_margin = Mm(20)
    sect.bottom_margin = Mm(20)

    disable_mirror_margins(doc)
    set_doc_defaults(doc)

    doc.add_paragraph(
        "Примечание к оформлению: документ подготовлен для односторонней печати "
        "(поля без зеркального расположения для переплёта)."
    ).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph()

    add_heading_center(doc, "2 Проектирование web-приложения BuildFlow", level=1)

    # --- 2.1 ---
    doc.add_heading("2.1 Описание структуры приложения", level=2)
    add_para(
        doc,
        "Целью разработки является web-приложение BuildFlow для учёта и сопровождения "
        "строительных проектов: планирование этапов и задач, обмен сообщениями внутри проекта, "
        "прикрепление фотофиксации выполнения работ, уведомления участников через Telegram, "
        "а также контроль статусов проекта и согласования планов работ заказчиком. Система "
        "ориентирована на асинхронное взаимодействие клиента и сервера через REST API и "
        "дополнительный канал реального времени на базе Socket.IO для оперативного обновления "
        "данных без полной перезагрузки интерфейса.",
    )
    add_para(
        doc,
        "Серверная часть реализована на платформе Node.js (образ контейнера node:20-alpine) "
        "с использованием фреймворка Express версии ^5.2.1. Доступ к данным выполняется через "
        "ORM Sequelize ^6.37.8 и драйвер tedious ^19.2.1 при работе с Microsoft SQL Server "
        "(контейнер mcr.microsoft.com/mssql/server:2019-latest). Для аутентификации "
        "используются JWT (jsonwebtoken ^9.0.3) и хеширование паролей bcrypt ^6.0.0; загрузка "
        "файлов организована модулем multer ^2.1.1. Интеграция с Telegram реализована библиотекой "
        "node-telegram-bot-api ^0.67.0.",
    )
    add_para(
        doc,
        "Клиентская часть построена как SPA на React ^19.2.4 с маршрутизацией react-router-dom "
        "^7.13.2 и HTTP-клиентом axios ^1.14.0; для синхронизации в реальном времени применяется "
        "socket.io-client ^4.8.3 (совместимо с серверным Socket.IO ^4.8.3). Пользовательский "
        "интерфейс формируется набором страниц и компонентов React с учётом ролевой модели "
        "(Гость, Заказчик, Прораб, Администратор).",
    )
    add_para(
        doc,
        "Production-сборка фронтенда упаковывается в статические файлы и обслуживается контейнером "
        "nginx:alpine. Конфигурация nginx выполняет раздачу SPA, перевод HTTP→HTTPS на порту 443, "
        "проксирование префиксов /api/, /uploads/ и /socket.io/ на сервис backend:5000 внутри "
        "Docker-сети compose-проекта. Это обеспечивает единую точку входа для браузера и "
        "корректную работу WebSocket-upgrade через reverse-proxy.",
    )
    add_para(
        doc,
        "Для воспроизводимости окружения используется Docker Compose: отдельные контейнеры для "
        "СУБД (buildflow_db), API-сервера (buildflow_backend) и frontend/nginx "
        "(buildflow_frontend_nginx), именованный том sql_data для данных MSSQL и привязки томов "
        "для каталога загрузок ./backend/uploads:/app/uploads и TLS-сертификатов ./certs:/etc/nginx/certs.",
    )

    # --- 2.2 ---
    doc.add_heading("2.2 Описание архитектуры web-приложения", level=2)
    add_para(
        doc,
        "Архитектура приложения многослойная и разделяет уровень представления (браузер + React SPA), "
        "уровень обратного прокси и статики (Nginx), уровень прикладной логики и REST/WebSocket "
        "(Node.js + Express + Socket.IO) и уровень хранения данных (Microsoft SQL Server). Такое "
        "разделение упрощает масштабирование и обслуживание компонентов.",
    )
    add_para(
        doc,
        "Уровень представления выполняется в web-браузере пользователя. Клиентское приложение "
        "получает статические ресурсы по HTTPS с контейнера frontend/nginx, выполняет запросы к "
        "REST API по маршруту /api/, загружает файлы через /uploads/, а канал Socket.IO использует "
        "тот же хост и порт 443; соединение устанавливается к nginx, который пробрасывает "
        "WebSocket-трафик на backend по пути /socket.io/.",
    )
    add_para(
        doc,
        "Уровень приложений размещается в контейнере backend и принимает обращения изнутри сети "
        "Compose по имени сервиса backend и порту 5000. Запросы маршрутизируются Express-маршрутами "
        "к контроллерам (авторизация, проекты, задачи и этапы, сервисы-справочники, администрирование "
        "пользователей). Middleware проверяет JWT и роли до выполнения бизнес-операций.",
    )
    add_para(
        doc,
        "Уровень данных реализован СУБД Microsoft SQL Server в контейнере db; приложение подключается "
        "к экземпляру по имени хоста db и порту 1433 с использованием параметров окружения DB_SERVER, "
        "DB_USER, DB_PASSWORD, DB_NAME из docker-compose.yml.",
    )
    add_para(
        doc,
        "Диаграмма развёртывания web-приложения BuildFlow должна быть представлена на рисунке 2.1 "
        "(рекомендуется использовать актуальный файл diagrams.net «диаграмма резвертывания.drawio» "
        "и вставить экспорт в PNG/SVG в документ пояснительной записки).",
    )
    p_fig = doc.add_paragraph()
    p_fig.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_fig.add_run("Рисунок 2.1 – Диаграмма развёртывания web-приложения BuildFlow")
    r.bold = True

    add_para(
        doc,
        "Логическая структура базы данных приводится в подразделе 2.3; для записки её следует также "
        "визуализировать ER-диаграммой (рисунок 2.2) по фактической схеме таблиц.",
    )
    p_fig2 = doc.add_paragraph()
    p_fig2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p_fig2.add_run("Рисунок 2.2 – Логическая схема базы данных BuildFlow")
    r2.bold = True

    # --- 2.3 ---
    doc.add_heading("2.3 Описание базы данных", level=2)
    add_para(
        doc,
        "Для хранения информации используется реляционная СУБД Microsoft SQL Server. Логическая модель "
        "отражает предметную область строительного проектирования и исполнения: справочник ролей и "
        "пользователей, каталог типовых услуг, проекты с составом участников, иерархию этапов и задач "
        "с возможными вложениями (файлы), а также сообщения проекта. Связи поддерживаются внешними "
        "ключами; для части коллекций заданы каскадные удаления на уровне ORM (например, этапы и "
        "сообщения при удалении проекта).",
    )

    add_table_caption(doc, "Таблица 2.1 – Структура таблицы Roles")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["id", "INT", "PRIMARY KEY, IDENTITY", "Уникальный идентификатор роли"],
            ["name", "NVARCHAR(255)", "NOT NULL, UNIQUE", "Наименование роли (Гость, Заказчик, Прораб, Администратор)"],
        ],
    )

    add_table_caption(doc, "Таблица 2.2 – Структура таблицы Users")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["id", "INT", "PRIMARY KEY, IDENTITY", "Уникальный идентификатор пользователя"],
            ["email", "NVARCHAR(255)", "NOT NULL, UNIQUE", "Электронная почта (логин)"],
            ["passwordHash", "NVARCHAR(255)", "NOT NULL", "Хеш пароля"],
            ["fullName", "NVARCHAR(255)", "NOT NULL", "ФИО пользователя"],
            ["status", "NVARCHAR(255)", "DEFAULT 'active'", "Статус учётной записи (active/blocked/deleted)"],
            ["telegramId", "NVARCHAR(255)", "NULL", "Идентификатор Telegram для уведомлений"],
            ["tgSettings", "NVARCHAR(MAX)", "DEFAULT JSON-строка", "Настройки уведомлений Telegram"],
            ["roleId", "INT", "FK → Roles.id", "Роль пользователя"],
            ["createdAt", "DATETIME2", "NOT NULL", "Дата создания записи"],
            ["updatedAt", "DATETIME2", "NOT NULL", "Дата обновления записи"],
        ],
    )

    add_table_caption(doc, "Таблица 2.3 – Структура таблицы Services")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["id", "INT", "PRIMARY KEY, IDENTITY", "Уникальный идентификатор услуги"],
            ["name", "NVARCHAR(255)", "NOT NULL", "Наименование услуги"],
            ["description", "NVARCHAR(MAX)", "NULL", "Описание услуги"],
        ],
    )

    add_table_caption(doc, "Таблица 2.4 – Структура таблицы Projects")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["id", "INT", "PRIMARY KEY, IDENTITY", "Уникальный идентификатор проекта"],
            ["name", "NVARCHAR(150)", "NOT NULL", "Название проекта"],
            ["description", "NVARCHAR(MAX)", "NULL", "Описание проекта"],
            ["startDate", "DATETIME2", "NULL", "Планируемая дата начала"],
            ["plannedEndDate", "DATETIME2", "NULL", "Планируемая дата окончания"],
            ["actualEndDate", "DATETIME2", "NULL", "Фактическая дата завершения"],
            ["status", "NVARCHAR(255)", "DEFAULT 'active'", "Статус проекта (active/completed/suspended)"],
            ["planStatus", "NVARCHAR(255)", "DEFAULT 'draft'", "Статус плана (draft/pending_approval/approved/rejected)"],
            ["createdAt", "DATETIME2", "NOT NULL", "Дата создания записи"],
            ["updatedAt", "DATETIME2", "NOT NULL", "Дата обновления записи"],
        ],
    )

    add_table_caption(doc, "Таблица 2.5 – Структура связующей таблицы ProjectUsers")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["projectId", "INT", "FK → Projects.id, часть составного PK", "Проект"],
            ["userId", "INT", "FK → Users.id, часть составного PK", "Участник проекта"],
        ],
    )

    add_table_caption(doc, "Таблица 2.6 – Структура таблицы ProjectStages")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["id", "INT", "PRIMARY KEY, IDENTITY", "Уникальный идентификатор этапа"],
            ["projectId", "INT", "FK → Projects.id", "Родительский проект"],
            ["name", "NVARCHAR(255)", "NOT NULL", "Наименование этапа"],
            ["description", "NVARCHAR(MAX)", "NULL", "Описание этапа"],
            ["startDate", "DATETIME2", "NULL", "Дата начала этапа"],
            ["plannedEndDate", "DATETIME2", "NULL", "Планируемое окончание этапа"],
            ["actualEndDate", "DATETIME2", "NULL", "Фактическое окончание этапа"],
            ["status", "NVARCHAR(255)", "DEFAULT 'в работе'", "Статус этапа (в работе / утверждено)"],
            ["order", "INT", "DEFAULT 0", "Порядок сортировки этапов"],
            ["createdAt", "DATETIME2", "NOT NULL", "Дата создания записи"],
            ["updatedAt", "DATETIME2", "NOT NULL", "Дата обновления записи"],
        ],
    )

    add_table_caption(doc, "Таблица 2.7 – Структура таблицы Tasks")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["id", "INT", "PRIMARY KEY, IDENTITY", "Уникальный идентификатор задачи"],
            ["stageId", "INT", "FK → ProjectStages.id", "Этап, к которому относится задача"],
            ["assignedUserId", "INT", "FK → Users.id", "Исполнитель задачи"],
            ["transferToUserId", "INT", "FK → Users.id, NULL", "Целевой исполнитель при передаче задачи"],
            ["description", "NVARCHAR(MAX)", "NOT NULL", "Описание задачи"],
            ["status", "NVARCHAR(255)", "DEFAULT 'новая'", "Статус (новая / в работе / выполнена)"],
            ["reportText", "NVARCHAR(MAX)", "NULL", "Текст отчёта о выполнении"],
            ["order", "INT", "DEFAULT 0", "Порядок сортировки внутри этапа"],
            ["createdAt", "DATETIME2", "NOT NULL", "Дата создания записи"],
            ["updatedAt", "DATETIME2", "NOT NULL", "Дата обновления записи"],
        ],
    )

    add_table_caption(doc, "Таблица 2.8 – Структура таблицы Attachments")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["id", "INT", "PRIMARY KEY, IDENTITY", "Уникальный идентификатор вложения"],
            ["taskId", "INT", "FK → Tasks.id", "Задача-владелец файла"],
            ["uploadedByUserId", "INT", "FK → Users.id", "Кто загрузил файл"],
            ["filePath", "NVARCHAR(255)", "NOT NULL", "Путь к файлу на сервере"],
            ["originalName", "NVARCHAR(255)", "NOT NULL", "Исходное имя файла"],
            ["createdAt", "DATETIME2", "NOT NULL", "Время загрузки"],
        ],
    )

    add_table_caption(doc, "Таблица 2.9 – Структура таблицы Messages")
    fill_table(
        doc,
        ["Столбец", "Тип данных", "Ограничение", "Назначение"],
        [
            ["id", "INT", "PRIMARY KEY, IDENTITY", "Уникальный идентификатор сообщения"],
            ["projectId", "INT", "FK → Projects.id", "Проект чата"],
            ["senderId", "INT", "FK → Users.id", "Отправитель"],
            ["text", "NVARCHAR(MAX)", "NOT NULL", "Текст сообщения"],
            ["isRead", "BIT", "DEFAULT 0", "Признак прочтения"],
            ["createdAt", "DATETIME2", "NOT NULL", "Время отправки"],
            ["updatedAt", "DATETIME2", "NOT NULL", "Время обновления записи"],
        ],
    )

    add_para(
        doc,
        "Основные связи типа «один ко многим»: Roles → Users; Projects → ProjectStages → Tasks; "
        "Users → Tasks (исполнение и запрос передачи); Tasks → Attachments; Projects → Messages; "
        "Users → Messages. Связь «многие ко многим» между Projects и Users реализована через таблицу "
        "ProjectUsers.",
    )

    # --- 2.4 ---
    doc.add_heading("2.4 Описание ролей пользователей и их функций", level=2)
    add_para(
        doc,
        "Ролевая модель обеспечивает разграничение доступа к операциям API и соответствующим разделам "
        "интерфейса. Роли задаются записью в таблице Roles и полем roleId пользователя; сервер дополнительно "
        "проверяет соответствие роли данным JWT.",
    )

    add_table_caption(doc, "Таблица 2.10 – Роли пользователей BuildFlow")
    fill_table(
        doc,
        ["Роль", "Общее описание", "Типичный доступ к данным"],
        [
            [
                "Гость",
                "Роль из справочника Roles (инициализируется при заполнении БД); может использоваться для ограниченного режима доступа при необходимости.",
                "Объём доступа задаётся политикой приложения; самостоятельная регистрация через API доступна только для ролей «Заказчик» и «Прораб».",
            ],
            [
                "Заказчик",
                "Заказчик работ: согласование плана, приёмка этапов, завершение проекта, отклонение задач при необходимости.",
                "Свои проекты и связанные с ними процессы согласования; операции approve/reject для этапов по маршрутам API.",
            ],
            [
                "Прораб",
                "Исполнитель: ведёт этапы и задачи, загружает отчёты и вложения, меняет порядок задач/этапов, запрашивает передачу задач.",
                "Назначенные задачи и проекты участия; операции CRUD задач в рамках полномочий, drag-and-drop порядка.",
            ],
            [
                "Администратор",
                "Управление пользователями и справочником услуг, создание/редактирование/удаление проектов, переазначение задач.",
                "Полный административный контур: /api/admin/*, управление сервисами, проекты без ограничения planStatus при удалении.",
            ],
        ],
    )

    add_table_caption(doc, "Таблица 2.11 – Функции роли «Администратор»")
    fill_table(
        doc,
        ["Функция", "Описание (по серверным маршрутам приложения)"],
        [
            ["Управление пользователями", "Просмотр списка, создание, изменение данных, удаление, смена статуса учётной записи (маршруты /api/admin/users)."],
            ["Управление справочником услуг", "CRUD записей Services (GET/POST/PUT/DELETE /api/services с проверкой роли)."],
            ["Управление проектами", "Создание, изменение и удаление проектов независимо от planStatus (маршруты /api/projects)."],
            ["Управление этапами и задачами", "Создание этапов, переименование, удаление вложений, создание задач, редактирование и удаление задач, переупорядочивание этапов и задач."],
        ],
    )

    add_table_caption(doc, "Таблица 2.12 – Функции роли «Прораб»")
    fill_table(
        doc,
        ["Функция", "Описание"],
        [
            ["Работа с задачами", "Просмотр назначенных задач (/api/tasks/my), обновление статуса и загрузка до 10 фото, текст отчёта; отклонение задачи заказчиком/прорабом по сценарию API."],
            ["Этапы", "Создание и удаление этапов, запрос/отмена согласования этапа, переименование (совместно с администратором)."],
            ["Передача задач", "Запрос передачи другому исполнителю, принятие/отклонение/отмена передачи."],
            ["Вложения", "Формирование списка вложений задачи через загрузку файлов; удаление вложения в рамках прав."],
            ["Порядок работ", "Изменение порядка этапов и задач (маршруты /api/order/...)."],
        ],
    )

    add_table_caption(doc, "Таблица 2.13 – Функции роли «Заказчик»")
    fill_table(
        doc,
        ["Функция", "Описание"],
        [
            ["Согласование плана", "Изменение plan-status проекта по API /api/projects/:id/plan-status."],
            ["Приёмка этапов", "Утверждение или отклонение этапа после запроса прораба."],
            ["Завершение проекта", "Фиксация завершения проекта (/api/projects/:id/complete)."],
            ["Коммуникации", "Чтение и отправка сообщений проекта, отметка прочитанным."],
        ],
    )

    add_table_caption(doc, "Таблица 2.14 – Функции роли «Гость» и общие операции")
    fill_table(
        doc,
        ["Функция", "Описание"],
        [
            ["Регистрация и вход", "Регистрация POST /api/auth/register (допускаются роли Заказчик и Прораб) и вход POST /api/auth/login; профиль текущего пользователя GET /api/auth/me с JWT."],
            ["Профиль", "Изменение данных и пароля PUT /api/auth/profile; отключение Telegram DELETE /api/auth/telegram."],
            ["Уведомления Telegram", "Привязка telegramId и рассылка событий по проектам средствами backend-сервиса (внешний HTTPS вызов Telegram Bot API)."],
        ],
    )

    add_para(
        doc,
        "Полный перечень пользовательских сценариев дополнительно отображается на диаграмме вариантов использования "
        "(UML Use Case) в приложении к пояснительной записке.",
    )

    # --- 2.5 ---
    doc.add_heading("2.5 Выводы по разделу", level=2)
    add_para(
        doc,
        "В результате проектирования определена целостная архитектура BuildFlow как SPA-приложения с выделенным "
        "reverse-proxy на nginx, REST API и WebSocket на Node.js/Express/Socket.IO и реляционным хранилищем "
        "Microsoft SQL Server. Контейнеризация через Docker Compose обеспечивает повторяемость конфигурации "
        "и совпадение параметров разработки и эксплуатации.",
    )
    add_para(
        doc,
        "Спроектированная база данных из девяти взаимосвязанных таблиц (включая связующую таблицу участников проектов ProjectUsers) покрывает "
        "ключевые сущности предметной области (пользователи и роли, услуги, проекты, этапы, задачи, вложения, сообщения) "
        "и поддерживает ограничение целостности связей.",
    )
    add_para(
        doc,
        "Ролевая модель (Гость, Заказчик, Прораб, Администратор) согласована с реализованными маршрутами API и "
        "JWT-авторизацией; это задаёт основу для безопасной реализации бизнес-процессов на следующих этапах "
        "разработки и тестирования приложения.",
    )

    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
