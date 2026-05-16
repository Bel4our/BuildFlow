# -*- coding: utf-8 -*-
"""Раздел 3 с расширенным набором листингов из кодовой базы BuildFlow."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

OUT = Path(r"C:\Users\Даниил\Downloads\Telegram Desktop\BuildFlow_section3_development.docx")

CODE = "Courier New"
CODE_SZ = Pt(10)


LISTING_DB = '''import { Sequelize } from 'sequelize';
import dotenv from 'dotenv';

dotenv.config();

const sequelize = new Sequelize(
  process.env.DB_NAME,
  process.env.DB_USER,
  process.env.DB_PASSWORD,
  {
    host: process.env.DB_SERVER,
    dialect: 'mssql',
    dialectOptions: {
      options: {
        encrypt: false,
        trustServerCertificate: true
      }
    },
    logging: false,
  }
);

export default sequelize;'''

LISTING_INDEX = '''app.use('/api/auth', authRoutes);
app.use('/api/projects', projectRoutes);
app.use('/api/tasks', taskRoutes);
app.use('/api/admin', adminRoutes);
app.use('/api/services', serviceRoutes);
app.use('/api/order', orderRoutes);

const startServer = async () => {
  while (true) {
    try {
      const masterSequelize = new Sequelize('master', process.env.DB_USER, process.env.DB_PASSWORD, {
        host: process.env.DB_SERVER,
        dialect: 'mssql',
        logging: false,
        dialectOptions: { options: { encrypt: false, trustServerCertificate: true } }
      });
      await masterSequelize.authenticate();
      await masterSequelize.query(
        `IF NOT EXISTS (SELECT * FROM sys.databases WHERE name = '${process.env.DB_NAME}')
         CREATE DATABASE [${process.env.DB_NAME}];`
      );
      await masterSequelize.close();

      await sequelize.authenticate();
      await sequelize.sync();
      await seedDatabase();

      server.listen(process.env.PORT || 5000, '0.0.0.0', () => {
        console.log(`СЕРВЕР РАБОТАЕТ НА ПОРТУ ${process.env.PORT || 5000}`);
      });
      break;
    } catch (err) {
      console.error('Ошибка:', err.message);
      await new Promise(r => setTimeout(r, 5000));
    }
  }
};'''

LISTING_SOCKET_IO_INIT = '''const server = http.createServer(app);

const io = new Server(server, {
  cors: {
    origin: "*",
    methods: ["GET", "POST"]
  }
});
initializeSocket(io);

app.use((req, res, next) => {
  req.io = io;
  next();
});'''

LISTING_AUTH_MW = '''export const verifyToken = (req, res, next) => {
  const token = req.headers['authorization']?.split(' ')[1];
  if (!token) return res.status(403).json({ message: 'Токен не предоставлен' });

  jwt.verify(token, process.env.JWT_SECRET, async (err, decoded) => {
    if (err) return res.status(401).json({ message: 'Неверный токен' });

    try {
      const user = await User.findByPk(decoded.id, { include: Role });
      if (!user || user.status === 'blocked' || user.status === 'deleted'
          || user.Role.name !== decoded.role) {
        return res.status(401).json({ message: 'Требуется авторизация' });
      }
      req.user = decoded;
      next();
    } catch (e) {
      return res.status(500).json({ message: 'Ошибка сервера' });
    }
  });
};

export const checkRole = (roles) => {
  return (req, res, next) => {
    if (!roles.includes(req.user.role)) {
      return res.status(403).json({ message: 'Нет доступа для вашей роли' });
    }
    next();
  };
};'''

LISTING_MULTER = '''const storage = multer.diskStorage({
  destination: (req, file, cb) => {
    const dir = 'uploads/';
    if (!fs.existsSync(dir)) {
      fs.mkdirSync(dir);
    }
    cb(null, dir);
  },
  filename: (req, file, cb) => {
    const uniqueSuffix = Date.now() + '-' + Math.round(Math.random() * 1E9);
    cb(null, uniqueSuffix + path.extname(file.originalname));
  }
});

const upload = multer({
  storage,
  fileFilter: (req, file, cb) => { cb(null, true); },
  limits: { fileSize: 10 * 1024 * 1024 }
});'''

LISTING_SOCKET_SERVICE = '''const initializeSocket = (io) => {
  io.on('connection', (socket) => {

    socket.on('join_project', (projectId) => {
      socket.join(projectId);
    });

    socket.on('leave_project', (projectId) => {
      socket.leave(projectId);
    });
  });
};

export default initializeSocket;'''

LISTING_MODELS_ASSOC = '''const ProjectUsers = sequelize.define('ProjectUsers', {}, { timestamps: false });
Project.belongsToMany(User, { as: 'Users', through: ProjectUsers, foreignKey: 'projectId' });
User.belongsToMany(Project, { as: 'Projects', through: ProjectUsers, foreignKey: 'userId' });

Project.hasMany(ProjectStage, { foreignKey: 'projectId', onDelete: 'CASCADE' });
ProjectStage.hasMany(Task, { foreignKey: 'stageId', onDelete: 'CASCADE' });

User.hasMany(Task, { foreignKey: 'assignedUserId' });
Task.belongsTo(User, { foreignKey: 'assignedUserId', as: 'worker' });

Task.hasMany(Attachment, { foreignKey: 'taskId', onDelete: 'CASCADE' });
Project.hasMany(Message, { foreignKey: 'projectId', onDelete: 'CASCADE' });'''

LISTING_REGISTER = '''export const register = async (req, res) => {
  try {
    const { email, password, fullName, roleName } = req.body;
    if (!email || !password || !fullName || !roleName)
      return res.status(400).json({ message: 'Все поля обязательны' });
    if (!['Заказчик', 'Прораб'].includes(roleName))
      return res.status(403).json({ message: 'Недопустимая роль' });
    if (!passRegex.test(password))
      return res.status(400).json({
        message: 'Пароль: мин. 8 символов, 1 буква, 1 цифра, 1 спецсимвол'
      });

    const existingUser = await User.findOne({ where: { email } });
    if (existingUser)
      return res.status(400).json({ message: 'Email уже существует' });

    const role = await Role.findOne({ where: { name: roleName } });
    if (!role)
      return res.status(500).json({ message: 'Роль не найдена' });

    const hashPassword = await bcrypt.hash(password, 10);
    await User.create({
      email,
      passwordHash: hashPassword,
      fullName: fullName.substring(0,100),
      roleId: role.id
    });
    res.status(201).json({ message: 'Регистрация успешна' });
  } catch (error) {
    res.status(500).json({ message: 'Ошибка сервера' });
  }
};'''

LISTING_LOGIN = '''const token = jwt.sign(
  {
    id: user.id,
    email: user.email,
    role: user.Role.name,
    fullName: user.fullName,
    hasTelegram: !!user.telegramId
  },
  process.env.JWT_SECRET,
  { expiresIn: '24h' }
);
res.json({ token, user: { id: user.id, role: user.Role.name, fullName: user.fullName } });'''

LISTING_CREATE_PROJECT = '''export const createProject = async (req, res) => {
  try {
    const { name, description, startDate, plannedEndDate, userIds } = req.body;
    if (new Date(plannedEndDate) < new Date(startDate))
      return res.status(400).json({ message: 'Dates error' });

    const users = await User.findAll({ where: { id: userIds }, include: [Role] });
    const clients = users.filter(u => u.Role.name === ROLES.CLIENT);
    const builders = users.filter(u => u.Role.name === ROLES.BUILDER);
    if (!clients.length || !builders.length || clients.length > 10 || builders.length > 10)
      return res.status(400).json({ message: 'Limits error (1-10)' });

    const project = await Project.create({ name, description, startDate, plannedEndDate });
    await project.setUsers(userIds);

    broadcastToProject(project.id, [ROLES.CLIENT, ROLES.BUILDER],
      `Вы назначены на новый проект: "${project.name}"`);
    res.status(201).json(project);
  } catch (error) {
    res.status(500).json({ message: 'Error' });
  }
};'''

LISTING_PLAN_STATUS = '''if (planStatus === PLAN_STATUSES.PENDING) {
  const projectWithStages = await Project.findByPk(projectId, {
    include: [{ model: ProjectStage, include: [Task] }]
  });
  for (const stage of projectWithStages.ProjectStages) {
    if (stage.Tasks.length === 0) {
      return res.status(400).json({
        message: `Этап "${stage.name}" не может быть пустым. Добавьте в него хотя бы одну задачу.`
      });
    }
    if (stage.Tasks.length > 50) {
      return res.status(400).json({
        message: `В этапе "${stage.name}" слишком много задач (максимум 50).`
      });
    }
  }
}
await project.update({ planStatus });
req.io.to(projectId.toString()).emit('stage_status_updated');'''

LISTING_CREATE_STAGE = '''const extractDateString = (val) => {
  if (!val) return null;
  if (/^\\d{4}-\\d{2}-\\d{2}$/.test(val)) return val;
  if (typeof val === 'string' && val.includes('T')) return val.split('T')[0];
  const d = new Date(val);
  if (isNaN(d.getTime())) return null;
  const year = d.getFullYear();
  const month = String(d.getMonth() + 1).padStart(2, '0');
  const day = String(d.getDate()).padStart(2, '0');
  return `${year}-${month}-${day}`;
};

export const createStage = async (req, res) => {
  try {
    const { projectId, name, description, startDate, plannedEndDate } = req.body;

    if (!name || name.trim().length === 0) {
      return res.status(400).json({ message: 'Название этапа не может быть пустым' });
    }

    const sdStr = extractDateString(startDate);
    const edStr = extractDateString(plannedEndDate);

    if (!sdStr || !edStr) {
      return res.status(400).json({ message: 'Даты этапа обязательны и должны быть корректными' });
    }

    const p = await Project.findByPk(projectId);
    if (!p) return res.status(404).json({ message: 'Проект не найден' });

    const sd = new Date(sdStr);
    const ed = new Date(edStr);
    const psd = p.startDate ? new Date(p.startDate) : null;
    const ped = p.plannedEndDate ? new Date(p.plannedEndDate) : null;

    sd.setHours(0,0,0,0);
    ed.setHours(0,0,0,0);
    if (psd && ped && (sd < psd || ed > ped || sd > ed)) {
      return res.status(400).json({ message: 'Даты этапа выходят за рамки сроков проекта' });
    }

    const stage = await ProjectStage.create({
      projectId,
      name: name.substring(0, 255),
      description: description || '',
      startDate: sequelize.literal(`CAST('${sdStr}' AS DATETIME)`),
      plannedEndDate: sequelize.literal(`CAST('${edStr}' AS DATETIME)`)
    });

    req.io.to(projectId.toString()).emit('stage_status_updated');
    res.status(201).json(stage);
  } catch (error) {
    res.status(500).json({ message: error.parent?.message || error.message || 'Внутренняя ошибка сервера' });
  }
};'''

LISTING_UPDATE_TASK_FRAGMENT = '''export const updateTask = async (req, res) => {
  try {
    const { status, reportText } = req.body;
    const task = await Task.findByPk(req.params.id, {
      include: [{ model: ProjectStage, include: [Project] }]
    });

    const oldStatus = task.status;
    if (status) task.status = status;
    if (reportText !== undefined) task.reportText = reportText.substring(0, 4000);
    await task.save();

    req.io.to(task.ProjectStage.projectId.toString()).emit('stage_status_updated');

    if (status && status !== oldStatus) {
      if (status === 'выполнена') {
        broadcastToProject(task.ProjectStage.Project.id, ['Заказчик'],
          `✅ В проекте "${task.ProjectStage.Project.name}" выполнена задача:\\n"${task.description}"`);
      } else if (status === 'в работе' && oldStatus === 'выполнена') {
        broadcastToProject(task.ProjectStage.Project.id, ['Заказчик'],
          `🔄 В проекте "${task.ProjectStage.Project.name}" задача возвращена в работу:\\n"${task.description}"`);
      } else if (status === 'новая' && oldStatus === 'в работе') {
        broadcastToProject(task.ProjectStage.Project.id, ['Заказчик'],
          `🔙 В проекте "${task.ProjectStage.Project.name}" задача возвращена в "Новые":\\n"${task.description}"`);
      }
    }

    if (req.files && req.files.length > 0) {
      const attachments = req.files.map(file => ({
        taskId: req.params.id,
        filePath: `/uploads/${file.filename}`,
        originalName: file.originalname,
        uploadedByUserId: req.user.id
      }));
      await Attachment.bulkCreate(attachments);
    }
    res.json(await Task.findByPk(req.params.id, { include: Attachment }));
  } catch (error) {
    res.status(500).json({ message: error.parent?.message || error.message || 'Error' });
  }
};'''

LISTING_ORDER = '''export const updateTaskOrder = async (req, res) => {
  const { orderedTaskIds } = req.body;
  const { stageId } = req.params;
  const t = await sequelize.transaction();
  try {
    for (let i = 0; i < orderedTaskIds.length; i++) {
      await Task.update(
        { order: i },
        { where: { id: orderedTaskIds[i], stageId: stageId }, transaction: t }
      );
    }
    await t.commit();
    const stage = await ProjectStage.findByPk(stageId);
    if (stage) {
      await triggerReapproval(stage.projectId, 'изменен порядок задач');
      req.io.to(stage.projectId.toString()).emit('stage_status_updated');
    }
    res.json({ message: 'Порядок задач обновлен' });
  } catch (error) {
    await t.rollback();
    res.status(500).json({ message: 'Ошибка обновления порядка задач' });
  }
};'''

LISTING_TG_NOTIFY = '''export const sendNotification = (userObj, text, projectId = null, options = {}) => {
  if (!bot || !userObj || !userObj.telegramId) return;
  try {
    let settings = { global: true, mutedProjects: [] };
    if (userObj.tgSettings) {
      try { settings = JSON.parse(userObj.tgSettings); } catch(e){}
    }
    if (!settings.global) return;
    if (projectId && settings.mutedProjects.includes(parseInt(projectId))) return;
    bot.sendMessage(userObj.telegramId, text, options);
  } catch (e) {}
};

export const broadcastToProject = async (projectId, rolesToNotify, text, options = {}) => {
  try {
    const project = await Project.findByPk(projectId, {
      include: [{ model: User, as: 'Users', include: [Role] }]
    });
    if (!project) return;
    const recipients = new Map();
    project.Users.forEach(u => {
      if (rolesToNotify.includes(u.Role.name)) recipients.set(u.id, u);
    });
    recipients.forEach(userObj => sendNotification(userObj, text, projectId, options));
  } catch (error) {}
};'''

LISTING_CRON = '''export const checkOverdueStages = async () => {
  try {
    const today = new Date();
    today.setHours(0, 0, 0, 0);
    const yesterday = new Date(today);
    yesterday.setDate(today.getDate() - 1);

    const justOverdueStages = await ProjectStage.findAll({
      where: {
        status: { [Op.ne]: 'утверждено' },
        plannedEndDate: { [Op.gte]: yesterday, [Op.lt]: today }
      },
      include: [{
        model: Project,
        where: { status: { [Op.ne]: 'completed' } },
        include: [{ model: User, as: 'Users', include: [Role] }]
      }]
    });

    for (const stage of justOverdueStages) {
      const targetUsers = stage.Project.Users.filter(u =>
        u.Role.name === ROLES.CLIENT || u.Role.name === ROLES.BUILDER
      );
      targetUsers.forEach(u => {
        sendNotification(u, `⚠️ Этап "${stage.name}" в проекте "${stage.Project.name}" просрочен...`, stage.Project.id);
      });
    }
  } catch (e) {}
};'''

LISTING_AXIOS = '''import axios from 'axios';

const instance = axios.create({
  baseURL: '/api'
});

instance.interceptors.request.use((config) => {
  const token = localStorage.getItem('token');
  if (token) {
    config.headers.Authorization = `Bearer ${token}`;
  }
  return config;
}, (error) => Promise.reject(error));

instance.interceptors.response.use(
  (response) => response,
  (error) => {
    if (error.response && error.response.status === 401) {
      localStorage.removeItem('token');
      window.location.href = '/login';
    }
    return Promise.reject(error);
  }
);

export default instance;'''

LISTING_AUTH_CTX = '''useEffect(() => {
  const token = localStorage.getItem('token');
  if (token) {
    try {
      const decoded = jwtDecode(token);
      if (decoded.exp * 1000 < Date.now()) {
        localStorage.removeItem('token');
        setUser(null);
      } else {
        setUser(decoded);
      }
    } catch (e) {
      localStorage.removeItem('token');
    }
  }
  setLoading(false);
}, []);

const login = (token) => {
  localStorage.setItem('token', token);
  setUser(jwtDecode(token));
};'''

LISTING_PROTECTED = '''const ProtectedRoute = ({ children, allowedRoles }) => {
  const { user } = useContext(AuthContext);

  if (!user) return <Navigate to="/login" />;
  if (allowedRoles && !allowedRoles.includes(user.role)) {
    return <div>Доступ запрещен! У вас нет прав для этой страницы.</div>;
  }
  return children;
};'''

LISTING_APP_ROUTES = '''<Routes>
  <Route path="/" element={<Home />} />
  <Route path="/login" element={<Login />} />
  <Route path="/register" element={<Register />} />
  <Route path="/project/:id" element={<ProtectedRoute><ProjectDetails /></ProtectedRoute>} />
  <Route path="/admin" element={<ProtectedRoute allowedRoles={['Администратор']}><AdminDashboard /></ProtectedRoute>} />
  <Route path="/builder" element={<ProtectedRoute allowedRoles={['Прораб']}><BuilderDashboard /></ProtectedRoute>} />
  <Route path="/my-tasks" element={<ProtectedRoute allowedRoles={['Прораб']}><MyTasks /></ProtectedRoute>} />
  <Route path="/client" element={<ProtectedRoute allowedRoles={['Заказчик']}><ClientDashboard /></ProtectedRoute>} />
  <Route path="/profile" element={<ProtectedRoute><Profile /></ProtectedRoute>} />
</Routes>'''


def set_margins_single_side(doc):
    s = doc.sections[0]
    s.left_margin = Mm(30)
    s.right_margin = Mm(15)
    s.top_margin = Mm(20)
    s.bottom_margin = Mm(20)


def disable_mirror(doc):
    el = doc.settings.element
    for child in list(el):
        if child.tag == qn("w:mirrorMargins"):
            el.remove(child)


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    if level <= 1:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.runs[0]
    r.bold = True
    p.paragraph_format.space_before = Pt(8)


def add_code_block(doc, code: str):
    for raw in code.split("\n"):
        line = raw if raw.strip() != "" else " "
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line)
        run.font.name = CODE
        run.font.size = CODE_SZ


def para(doc, txt, justify=True):
    p = doc.add_paragraph(txt)
    if justify:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def main():
    doc = Document()
    disable_mirror(doc)
    set_margins_single_side(doc)

    norm = doc.styles["Normal"]
    norm.font.name = "Times New Roman"
    norm.font.size = Pt(14)

    add_heading(doc, "3 Разработка web-приложения BuildFlow", 1)

    doc.add_heading("3.1 Программная платформа Node.js", 2)
    para(
        doc,
        "Серверная часть выполнена на Node.js (ES Modules). Используется HTTP-сервер из модуля `http`, поверх которого монтируются Express и Socket.IO. "
        "Ниже приведён фрагмент подключения маршрутов и циклическое ожидание готовности Microsoft SQL Server с автоматическим созданием базы при первом запуске.",
    )
    add_caption(doc, "Листинг 3.1 – Подключение REST-маршрутов и запуск после готовности БД (backend/index.js, фрагмент)")
    add_code_block(doc, LISTING_INDEX)

    para(
        doc,
        "Инициализация Socket.IO и передача экземпляра `io` в обработчики запросов через `req.io` показаны ниже.",
    )
    add_caption(doc, "Листинг 3.2 – Создание Socket.IO и внедрение в middleware (backend/index.js, фрагмент)")
    add_code_block(doc, LISTING_SOCKET_IO_INIT)

    doc.add_heading("3.2 Система управления базами данных и ORM Sequelize", 2)
    para(
        doc,
        "Доступ к Microsoft SQL Server настраивается в `backend/config/db.js`; диалект `mssql` использует драйвер tedious.",
    )
    add_caption(doc, "Листинг 3.3 – Подключение Sequelize к SQL Server (backend/config/db.js)")
    add_code_block(doc, LISTING_DB)

    add_caption(doc, "Листинг 3.4 – Основные ассоциации Sequelize (backend/models/index.js, фрагмент)")
    add_code_block(doc, LISTING_MODELS_ASSOC)

    doc.add_heading("3.3 Безопасность: JWT, роли и загрузка файлов", 2)
    add_caption(doc, "Листинг 3.5 – Проверка JWT и роли (backend/middleware/authMiddleware.js)")
    add_code_block(doc, LISTING_AUTH_MW)

    add_caption(
        doc,
        "Листинг 3.6 – Конфигурация multer: уникальные имена файлов и предел 10 МБ (backend/middleware/uploadMiddleware.js, фрагмент)",
    )
    add_code_block(doc, LISTING_MULTER)

    doc.add_heading("3.4 Режим реального времени и фоновые уведомления", 2)
    add_caption(doc, "Листинг 3.7 – Подключение клиента Socket.IO к комнатам проекта (backend/services/socketService.js)")
    add_code_block(doc, LISTING_SOCKET_SERVICE)

    doc.add_heading("3.5 Регистрация и выдача токена", 2)
    add_caption(doc, "Листинг 3.8 – Регистрация с ограничением ролей (backend/controllers/authController.js, фрагмент)")
    add_code_block(doc, LISTING_REGISTER)

    para(doc, "После успешной проверки пароля клиент получает JWT; фрагмент формирования ответа приведён ниже.")
    add_caption(doc, "Листинг 3.9 – Полезная нагрузка JWT и ответ клиенту (backend/controllers/authController.js, фрагмент)")
    add_code_block(doc, LISTING_LOGIN)

    doc.add_heading("3.6 Проекты и статус плана работ", 2)
    add_caption(doc, "Листинг 3.10 – Создание проекта и назначение участников (backend/controllers/projectController.js, фрагмент)")
    add_code_block(doc, LISTING_CREATE_PROJECT)

    para(
        doc,
        "При переводе плана на согласование проверяется, что каждый этап содержит от одной до пятидесяти задач; после обновления рассылается событие в комнату проекта.",
    )
    add_caption(doc, "Листинг 3.11 – Фрагмент согласования плана проекта (backend/controllers/projectController.js)")
    add_code_block(doc, LISTING_PLAN_STATUS)

    doc.add_heading("3.7 Этапы и задачи, вложения", 2)
    add_caption(
        doc,
        "Листинг 3.12 – Создание этапа с явным CAST дат под SQL Server (backend/controllers/taskController.js, упрощённый фрагмент)",
    )
    add_code_block(doc, LISTING_CREATE_STAGE)

    para(
        doc,
        "Ниже — обобщённый фрагмент обработчика обновления задачи: смена статуса, отчёт, рассылка `stage_status_updated`, создание вложений после `multer.array('photos', 10)`.",
    )
    add_caption(doc, "Листинг 3.13 – Обновление задачи и вложения (backend/controllers/taskController.js, фрагмент)")
    add_code_block(doc, LISTING_UPDATE_TASK_FRAGMENT)

    doc.add_heading("3.8 Транзакции порядка этапов и задач", 2)
    add_caption(doc, "Листинг 3.14 – Атомарное обновление порядка задач (backend/controllers/orderController.js, фрагмент)")
    add_code_block(doc, LISTING_ORDER)

    doc.add_heading("3.9 Интеграция Telegram Bot API", 2)
    add_caption(doc, "Листинг 3.15 – Отправка и групповая рассылка в Telegram с учётом настроек (backend/services/telegramBot.js, фрагмент)")
    add_code_block(doc, LISTING_TG_NOTIFY)

    doc.add_heading("3.10 Периодический контроль дедлайнов этапов", 2)
    add_caption(
        doc,
        "Листинг 3.16 – Ежесуточный поиск только что просроченных этапов (`setInterval` в index.js вызывает эту функцию) (backend/services/cronService.js, фрагмент)",
    )
    add_code_block(doc, LISTING_CRON)

    doc.add_heading("3.11 Клиент: Axios, контекст авторизации и маршруты", 2)
    add_caption(doc, "Листинг 3.17 – Экземпляр Axios с JWT и перенаправлением при 401 (frontend/src/api/axios.js)")
    add_code_block(doc, LISTING_AXIOS)

    add_caption(doc, "Листинг 3.18 – Инициализация пользователя из JWT при загрузке SPA (frontend/src/context/AuthContext.js, фрагмент)")
    add_code_block(doc, LISTING_AUTH_CTX)

    add_caption(doc, "Листинг 3.19 – Компонент защиты маршрутов по ролям (frontend/src/components/ProtectedRoute.js)")
    add_code_block(doc, LISTING_PROTECTED)

    add_caption(doc, "Листинг 3.20 – Маршрутизация приложения и привязка дашбордов к ролям (frontend/src/App.js, фрагмент)")
    add_code_block(doc, LISTING_APP_ROUTES)

    doc.add_heading("3.12 Выводы по разделу", 2)
    para(
        doc,
        "Раздел содержит листинги, непосредственно подтверждающие реализацию: конфигурация Sequelize под SQL Server, разделение REST и WebSocket, "
        "JWT и роли, загрузку файлов, бизнес-правила проекта и плана работ, транзакции при перестановке задач, интеграцию Telegram и ежесуточные напоминания, "
        "а также ключевые фрагменты React SPA (Axios interceptors и защищённые маршруты).",
    )

    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
